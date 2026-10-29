import re
from datetime import date
from pathlib import Path

from ..skills import extract_skills
from .llm import LLMProvider, LLMResult

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{8,}\d)")
MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
DATE_RANGE_RE = re.compile(
    rf"((?:{MONTHS})[a-z]*\.?\s+\d{{4}}|\d{{4}})\s*[-–—to]+\s*((?:{MONTHS})[a-z]*\.?\s+\d{{4}}|\d{{4}}|present|current|now|till date)",
    re.IGNORECASE,
)
HEADINGS = {
    "summary": ["summary", "profile", "objective", "about"],
    "skills": ["skills", "technical skills", "technologies", "tech stack"],
    "experience": ["experience", "work experience", "employment", "work history", "professional experience"],
    "education": ["education", "academic", "qualifications"],
    "projects": ["projects", "personal projects", "key projects"],
    "certifications": ["certifications", "certificates"],
}
DEGREE_RE = re.compile(
    r"\b(b\.?tech|m\.?tech|b\.?e\.?|m\.?e\.?|b\.?sc|m\.?sc|bca|mca|mba|bachelor|master|ph\.?d|diploma)\b", re.IGNORECASE
)


def extract_text(path: Path, content_type: str) -> str:
    if content_type == "application/pdf" or path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    from docx import Document

    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _heading_of(line: str) -> str | None:
    stripped = line.strip().strip(":").lower()
    if len(stripped) > 40:
        return None
    for section, names in HEADINGS.items():
        if stripped in names:
            return section
    return None


def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = _heading_of(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line.strip())
    return sections


def _parse_month_year(value: str, end: bool = False) -> date | None:
    value = value.strip().lower()
    if value in ("present", "current", "now", "till date"):
        return date.today()
    m = re.match(rf"({MONTHS})[a-z]*\.?\s+(\d{{4}})", value)
    if m:
        month = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(m.group(1)[:3]) + 1
        return date(int(m.group(2)), month, 1)
    if re.fullmatch(r"\d{4}", value):
        return date(int(value), 12 if end else 1, 1)
    return None


def experience_entries(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    for i, line in enumerate(lines):
        m = DATE_RANGE_RE.search(line)
        if not m:
            continue
        start = _parse_month_year(m.group(1))
        end = _parse_month_year(m.group(2), end=True)
        title_line = line[: m.start()].strip(" -–|,") or (lines[i - 1] if i > 0 else "")
        entries.append(
            {
                "title": title_line[:120],
                "start": start.isoformat() if start else m.group(1),
                "end": end.isoformat() if end else m.group(2),
                "text": line,
            }
        )
    return entries


def total_years(entries: list[dict]) -> float | None:
    spans = []
    for e in entries:
        try:
            start = date.fromisoformat(e["start"])
            end = date.fromisoformat(e["end"])
        except (ValueError, TypeError, KeyError):
            continue
        if end > start:
            spans.append((start, end))
    if not spans:
        return None
    spans.sort()
    merged = [spans[0]]
    for start, end in spans[1:]:
        last_start, last_end = merged[-1]
        if start <= last_end:
            merged[-1] = (last_start, max(last_end, end))
        else:
            merged.append((start, end))
    days = sum((end - start).days for start, end in merged)
    return round(days / 365.25, 1)


def parse_resume(text: str, llm: LLMProvider | None = None) -> tuple[dict, LLMResult | None]:
    sections = split_sections(text)
    header = sections.get("header", [])
    email = EMAIL_RE.search(text)
    phone = PHONE_RE.search(text)
    name = ""
    for line in header[:5]:
        if "@" in line or PHONE_RE.search(line) or len(line.split()) > 5:
            continue
        name = line.strip()
        break

    experience = experience_entries(sections.get("experience", []) or [l for l in text.splitlines() if DATE_RANGE_RE.search(l)])
    skills_text = " ".join(sections.get("skills", [])) or text
    parsed = {
        "name": name,
        "email": email.group(0) if email else "",
        "phone": phone.group(0).strip() if phone else "",
        "summary": " ".join(sections.get("summary", []))[:1000],
        "skills": extract_skills(skills_text) or extract_skills(text),
        "experience": experience,
        "experience_years": total_years(experience),
        "education": [{"text": l} for l in sections.get("education", []) if DEGREE_RE.search(l) or len(l) > 8][:10],
        "projects": [{"text": l} for l in sections.get("projects", [])][:20],
        "certifications": sections.get("certifications", [])[:10],
    }

    result = None
    if llm and llm.enabled:
        data, result = llm.complete_json(
            "You extract resume data. Use only what is in the resume text; never add skills or employers that are not there.",
            "Return JSON with keys name, summary, skills (list of strings), experience (list of {title, company, start, end, "
            "text}), education (list of {text}), projects (list of {text}).\n\nResume:\n" + text[:12000],
        )
        if data:
            for key in ("name", "summary", "skills", "experience", "education", "projects"):
                if data.get(key):
                    parsed[key] = data[key]
            parsed["experience_years"] = total_years(parsed["experience"]) or parsed["experience_years"]
    return parsed, result
